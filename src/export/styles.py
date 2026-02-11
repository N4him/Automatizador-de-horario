"""
Estilos para exportación Excel
"""
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side


# COLORES
COLOR_NARANJA = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
COLOR_VERDE = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
COLOR_AZUL = PatternFill(start_color="87CEEB", end_color="87CEEB", fill_type="solid")
COLOR_VACIO = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
COLOR_SIN_MONITOR = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
COLOR_HEADER = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
COLOR_TITULO_SALA = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
COLOR_MONITOR = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
COLOR_CLASE = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

# FUENTES
FUENTE_NEGRA = Font(bold=False, size=8, color="000000")
FUENTE_HEADER = Font(bold=True, size=9, color="000000")
FUENTE_TITULO = Font(bold=True, size=11, color="FFFFFF")
FUENTE_MONITOR = Font(bold=True, size=12, color="FFFFFF")
FUENTE_NORMAL = Font(size=8, color="000000")
FUENTE_ERROR = Font(size=7, color="FF0000", bold=True)

# ALINEACIÓN
ALINEACION_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

# BORDES
BORDE = Border(
    left=Side(style='thin', color='000000'),
    right=Side(style='thin', color='000000'),
    top=Side(style='thin', color='000000'),
    bottom=Side(style='thin', color='000000')
)